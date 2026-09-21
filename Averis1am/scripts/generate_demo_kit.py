#!/usr/bin/env python3
"""Build the live-demo kit: documents, emails to send, and a runbook.

Every scenario is sent by a person from a real mailbox to the connected SDOC
mailbox, so it has to survive the live path: mailbox query, classification,
preflight, extraction, the independent Gemini read, and case grouping.

    python scripts/generate_demo_kit.py                 # batch 1 -> ../demo-kit
    python scripts/generate_demo_kit.py --batch 2       # fresh references for a rehearsal
    python scripts/generate_demo_kit.py --check         # run every email through the pipeline
    python scripts/generate_demo_kit.py --check --verify   # ... including the Gemini second read

Shipment references are ``5PKG-<batch>26<nn>``. A reference that already has a
case in the store joins that case, so use a new batch for every run-through.
All companies, vessels and values are fictitious.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from reportlab.lib import colors, pdfencrypt  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402

CARRIER = "KESTREL OCEAN LINES"
FOOTER = "Fictitious document generated for the SDOC demo - not a real shipment"


# ---------------------------------------------------------------------------
# Documents
#
# Layout rules come from the PDF reader, not taste: the title must be the first
# line on the page, and every value sits on the same line as its bold label.
# Nothing else may share a line with a label, and no free text may contain a
# colon, or it would be read as a field.
# ---------------------------------------------------------------------------
def write_pdf(path, title, issuer, sections, password=None):
    encrypt = None
    if password:
        encrypt = pdfencrypt.StandardEncryption(password, ownerPassword=password + "-owner")
    pdf = canvas.Canvas(str(path), pagesize=A4, encrypt=encrypt)
    width, height = A4
    pdf.setTitle(title)
    pdf.setFillColor(colors.HexColor("#1c1c1c"))
    pdf.setFont("Helvetica-Bold", 17)
    pdf.drawString(54, height - 62, title)
    pdf.setFont("Helvetica", 9)
    pdf.setFillColor(colors.HexColor("#555555"))
    pdf.drawString(54, height - 80, issuer)
    pdf.setStrokeColor(colors.HexColor("#e8892a"))
    pdf.setLineWidth(2)
    pdf.line(54, height - 92, width - 54, height - 92)

    y = height - 124
    for heading, rows in sections:
        pdf.setFillColor(colors.HexColor("#8a5a2b"))
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(54, y, heading.upper())
        y -= 10
        for label, value in rows:
            pdf.setFillColor(colors.HexColor("#f4f2ee"))
            pdf.rect(54, y - 20, 150, 26, stroke=0, fill=1)
            pdf.setStrokeColor(colors.HexColor("#d9d6d0"))
            pdf.setLineWidth(0.6)
            pdf.rect(54, y - 20, width - 108, 26, stroke=1, fill=0)
            pdf.setFillColor(colors.HexColor("#1c1c1c"))
            pdf.setFont("Helvetica-Bold", 9.5)
            pdf.drawString(62, y - 11, label)
            pdf.setFont("Helvetica", 10)
            pdf.drawString(214, y - 11, value)
            y -= 26
        y -= 18
    pdf.setFont("Helvetica", 7.5)
    pdf.setFillColor(colors.HexColor("#888888"))
    pdf.drawString(54, 40, FOOTER)
    pdf.save()


def si_doc(ref, s, labels=None):
    lab = {"pol": "Port of Loading", "pod": "Port of Discharge",
           "count": "No. of Containers", "weight": "Gross Weight", **(labels or {})}
    return (
        "SHIPPING INSTRUCTION",
        f"Issued by {s['shipper'].title()} to {CARRIER.title()}",
        [
            ("Booking", [("Booking No.", ref), ("Vessel / Voyage", s["vessel"])]),
            ("Parties", [("Shipper", s["shipper"]), ("Consignee", s["consignee"]),
                         ("Notify Party", s["notify"])]),
            ("Routing", [(lab["pol"], s["pol"]), (lab["pod"], s["pod"])]),
            ("Cargo", [(lab["count"], s["count"]), (lab["weight"], s["weight"]),
                       ("Description of Goods", s["goods"]), ("HS Code", s["hs"])]),
        ],
    )


def bl_doc(ref, s, version=1):
    return (
        "BILL OF LADING - DRAFT FOR APPROVAL",
        f"{CARRIER} - draft {version} issued for shipper review",
        [
            ("Reference", [("B/L No.", f"KOL{ref.replace('-', '')}"), ("Booking No.", ref),
                           ("Vessel / Voyage", s["vessel"])]),
            ("Parties", [("Shipper", s["shipper"]), ("Consignee", s["consignee"]),
                         ("Notify Party", s["notify"])]),
            ("Routing", [("Port of Loading", s["pol"]), ("Port of Discharge", s["pod"])]),
            ("Cargo", [("Container Count", s["count"]), ("Gross Weight", s["weight"]),
                       ("Description of Goods", s["goods"]), ("Freight Terms", "PREPAID")]),
        ],
    )


def write_packing_list(path, ref, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Packing List"
    ws.append(["Booking No.", ref])
    ws.append([])
    ws.append(["Container", "Seal", "Packages", "Gross (kg)"])
    for row in rows:
        ws.append(list(row))
    for column in "ABCD":
        ws.column_dimensions[column].width = 18
    wb.save(path)


def write_booking_docx(path, ref):
    import docx
    document = docx.Document()
    document.add_heading("BOOKING CONFIRMATION", level=1)
    document.add_paragraph(f"{CARRIER.title()} confirms space for the booking below.")
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in [
        ("Booking No.", ref), ("Vessel / Voyage", "KESTREL HARMONY 118E"),
        ("Load Port", "PORT KLANG, MALAYSIA"), ("Destination", "JEBEL ALI, UAE"),
        ("Equipment", "2 x 40'HC"), ("Cargo cut-off", "Friday 16:00"),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.save(path)


# ---------------------------------------------------------------------------
# Scenarios
# ---------------------------------------------------------------------------
def scenarios(batch):
    ref = lambda n: f"5PKG-{batch}26{n:02d}"  # noqa: E731

    s1 = dict(shipper="BLUEWATER GLOVE INDUSTRIES SDN. BHD.", consignee="NORDKAP MEDICAL IMPORT GMBH",
              notify="NORDKAP LOGISTIK GMBH", pol="PORT KLANG, MALAYSIA", pod="HAMBURG, GERMANY",
              count="2 x 40'HC", weight="24,850 KG", goods="NITRILE EXAMINATION GLOVES, 1,960 CARTONS",
              hs="4015.19", vessel="KESTREL HARMONY 118W")
    s2 = dict(shipper="TUAH RUBBERWOOD FURNITURE SDN. BHD.", consignee="PACIFIC CREST DISTRIBUTION INC.",
              notify="CREST CUSTOMS BROKERAGE LLC", pol="PASIR GUDANG, MALAYSIA", pod="LONG BEACH, USA",
              count="3 x 40'HC", weight="51,300 KG", goods="RUBBERWOOD DINING SETS, KNOCK-DOWN, 612 PACKAGES",
              hs="9403.60", vessel="KESTREL VANGUARD 204E")
    s2_swapped = dict(s2, consignee=s2["notify"], notify=s2["consignee"])
    s3_si = dict(shipper="Sri Maju Oleochemicals Sdn Bhd", consignee="Harbor & Line Trading Pte. Ltd.",
                 notify="Harbor & Line Trading Pte. Ltd.", pol="Port Klang (MYPKG)", pod="Singapore (SGSIN)",
                 count="3 x 20' ISO tank", weight="61.2 MT", goods="Refined glycerine 99.7% USP, in bulk",
                 hs="1520.00", vessel="Kestrel Straits 031N")
    s3_bl = dict(s3_si, shipper="SRI MAJU OLEOCHEMICALS SDN. BHD.",
                 consignee="HARBOR AND LINE TRADING PTE LTD", notify="HARBOR AND LINE TRADING PTE LTD",
                 pol="PORT KLANG", pod="SINGAPORE", count="THREE (3) X 20'TK", weight="61,200.000 KGS",
                 goods="REFINED GLYCERINE 99.7% USP IN BULK", vessel="KESTREL STRAITS 031N")
    s4 = dict(shipper="TERAS ELECTRONICS SDN. BHD.", consignee="VANDERHOEK COMPONENTS B.V.",
              notify="VANDERHOEK COMPONENTS B.V.", pol="TANJUNG PELEPAS, MALAYSIA",
              pod="ROTTERDAM, NETHERLANDS", count="1 x 40'HC", weight="11,640 KG",
              goods="PRINTED CIRCUIT BOARD ASSEMBLIES, 38 PALLETS", hs="8534.00",
              vessel="KESTREL MERIDIAN 077W")
    s5 = dict(shipper="HIJAU PALM SPECIALTIES SDN. BHD.", consignee="OSAKA FINE FOODS CO., LTD.",
              notify="OSAKA FINE FOODS CO., LTD.", pol="PORT KLANG, MALAYSIA", pod="OSAKA, JAPAN",
              count="1 x 20'GP", weight="18,960 KG", goods="RED PALM OIL IN DRUMS, 80 DRUMS",
              hs="1511.90", vessel="KESTREL ORIENT 452N")

    return [
        {
            "key": "01-transposed-weight", "ref": ref(1),
            "title": "The transposed digit",
            "result": "DISCREPANCY on gross weight only (6 of 7 fields match)",
            "story": "The SI says 24,850 KG. The carrier's draft BL says 24,580 KG: two digits swapped. "
                     "Nothing else differs. This is the error a tired reviewer reads straight past, and a "
                     "wrong declared weight on a BL causes customs queries and amendment fees after "
                     "the vessel sails.",
            "show": ["Needs action list: the case appears under Discrepancy.",
                     "Open it: Gross weight is the only highlighted row.",
                     "Evidence pane: both crops show the printed value with a box around it.",
                     "Prepare message: SDOC drafts the correction request quoting both values. It never sends it."],
            "say": "Six fields match, one doesn't, and it shows you where on each page.",
            "files": lambda d: [
                pdf(d, f"{ref(1)}_SI.pdf", si_doc(ref(1), s1)),
                pdf(d, f"{ref(1)}_BL.pdf", bl_doc(ref(1), dict(s1, weight="24,580 KG"))),
            ],
            "emails": [{
                "subject": f"Draft BL check - {ref(1)} (Hamburg)",
                "body": f"Hi team,\n\nPlease compare the attached SI and draft BL for booking {ref(1)} "
                        "before we approve the draft with the carrier.\n\nThanks,\nAina\nExport Documentation",
                "attach": [f"{ref(1)}_SI.pdf", f"{ref(1)}_BL.pdf"],
                "expect": "DISCREPANCY",
            }],
        },
        {
            "key": "02-swapped-consignee", "ref": ref(2),
            "title": "Swapped parties, then the fix",
            "result": "DISCREPANCY on consignee and notify party, then VERIFIED after the corrected draft",
            "story": "The carrier's first draft swaps the consignee and the notify party. The consignee is "
                     "the party entitled to the cargo, so this is the error that stops a release at "
                     "destination. The carrier then sends a corrected draft, and the same case closes itself.",
            "show": ["After email 1: Discrepancy on Consignee and Notify party. The names are identical, just in the wrong rows.",
                     "After email 2: the same case turns Verified. It is not a new case.",
                     "Email and documents tab: BL v1 and v2 both listed, v2 active.",
                     "History tab: discrepancy found, corrected draft received, verified, open review task cancelled."],
            "say": "One shipment, one case: the correction closes it instead of opening a new ticket.",
            "files": lambda d: [
                pdf(d, f"{ref(2)}_SI.pdf", si_doc(ref(2), s2)),
                pdf(d, f"{ref(2)}_BL_v1.pdf", bl_doc(ref(2), s2_swapped, 1)),
                pdf(d, f"{ref(2)}_BL_v2.pdf", bl_doc(ref(2), s2, 2)),
            ],
            "emails": [
                {
                    "subject": f"Draft BL for approval - {ref(2)}",
                    "body": f"Hi team,\n\nFirst draft BL from the carrier for {ref(2)}. Please check it "
                            "against our SI.\n\nRegards,\nDaniel\nShipping Desk",
                    "attach": [f"{ref(2)}_SI.pdf", f"{ref(2)}_BL_v1.pdf"],
                    "expect": "DISCREPANCY",
                },
                {
                    "subject": f"Corrected draft BL - {ref(2)}",
                    "body": f"Hi team,\n\nThe carrier has corrected the draft BL for {ref(2)}. SI attached "
                            "again so you have both together. Please re-check.\n\nRegards,\nDaniel\nShipping Desk",
                    "attach": [f"{ref(2)}_SI.pdf", f"{ref(2)}_BL_v2.pdf"],
                    "expect": "VERIFIED",
                },
            ],
        },
        {
            "key": "03-different-formats-same-facts", "ref": ref(3),
            "title": "Different formats, same facts",
            "result": "VERIFIED: no false alarms. The Excel packing list is recognised and set aside",
            "story": "The shipper types its SI in its own style: tonnes instead of kilograms, port codes in "
                     "brackets, '&' and 'Pte. Ltd.', different field names. The carrier prints everything in "
                     "capitals with other labels. A plain text diff would flag almost every row. None of the "
                     "differences change a fact, so SDOC must not raise one.",
            "show": ["Case lands in Verified, not in Needs action.",
                     "Field comparison: '61.2 MT' vs '61,200.000 KGS', 'Port Klang (MYPKG)' vs 'PORT KLANG', "
                     "'Harbor & Line Trading Pte. Ltd.' vs 'HARBOR AND LINE TRADING PTE LTD', all Match.",
                     "Email and documents: three attachments. The packing list (Excel icon) is identified as a "
                     "supporting document and not compared."],
            "say": "Strict about facts, relaxed about formatting. False alarms are what make teams stop trusting a tool.",
            "files": lambda d: [
                pdf(d, f"{ref(3)}_SI.pdf", si_doc(ref(3), s3_si, labels={
                    "pol": "Load Port", "pod": "Discharge Port", "count": "Number of Containers",
                    "weight": "Total Gross Weight"})),
                pdf(d, f"{ref(3)}_BL.pdf", bl_doc(ref(3), s3_bl)),
                xlsx(d, f"{ref(3)}_Packing_List.xlsx", ref(3), [
                    ("KOLU 204118-6", "MY7781201", "1 ISO tank", 20400),
                    ("KOLU 204119-1", "MY7781202", "1 ISO tank", 20380),
                    ("KOLU 204120-4", "MY7781203", "1 ISO tank", 20420),
                ]),
            ],
            "emails": [{
                "subject": f"SI and draft BL comparison - {ref(3)} (Singapore)",
                "body": f"Hi team,\n\nPlease compare the SI and the draft BL for {ref(3)}. The packing list "
                        "is attached for reference only.\n\nBest,\nMei Ling\nOperations",
                "attach": [f"{ref(3)}_SI.pdf", f"{ref(3)}_BL.pdf", f"{ref(3)}_Packing_List.xlsx"],
                "expect": "VERIFIED",
            }],
        },
        {
            "key": "04-bl-to-follow", "ref": ref(4),
            "title": "The BL is still on its way",
            "result": "WAITING (no alarm, nothing for the team to do), then VERIFIED when both documents are in",
            "story": "Documents rarely arrive together. The shipper sends the SI first and the carrier's draft "
                     "follows hours later. SDOC recognises an incomplete set and waits. It does not raise an "
                     "error or put anything in the action queue. If the BL never arrives, the case escalates "
                     "on its own after the waiting window (24 hours as configured).",
            "show": ["After email 1: the case appears under All cases as Waiting, not under Needs action.",
                     "Case guidance: 'Waiting for the full document set. Nothing to do yet.'",
                     "Prepare message: a ready draft asking for the missing draft BL.",
                     "After email 2: the same case turns Verified."],
            "say": "It knows the difference between 'wrong' and 'not here yet'.",
            "files": lambda d: [
                pdf(d, f"{ref(4)}_SI.pdf", si_doc(ref(4), s4)),
                pdf(d, f"{ref(4)}_BL.pdf", bl_doc(ref(4), s4)),
            ],
            "emails": [
                {
                    "subject": f"Draft BL to follow - {ref(4)}",
                    "body": f"Hi team,\n\nOur SI for booking {ref(4)} is attached. The carrier will send the "
                            "draft BL later today, so please hold the check until it arrives.\n\n"
                            "Thanks,\nFaizal\nExport Documentation",
                    "attach": [f"{ref(4)}_SI.pdf"],
                    "expect": "WAITING",
                },
                {
                    "subject": f"Draft BL received - {ref(4)}",
                    "body": f"Hi team,\n\nThe carrier's draft BL for {ref(4)} is attached, with the SI again so "
                            "both are together. Please go ahead with the check.\n\nThanks,\nFaizal\nExport Documentation",
                    "attach": [f"{ref(4)}_SI.pdf", f"{ref(4)}_BL.pdf"],
                    "expect": "VERIFIED",
                },
            ],
        },
        {
            "key": "05-locked-bl", "ref": ref(5),
            "title": "A password-protected draft BL",
            "result": "BLOCKED: attachment_encrypted, stopped before any AI reads it",
            "story": "Carriers and banks often password-protect documents. SDOC cannot read a locked file, so "
                     "it must not guess. The safety check stops the file at the door (size, true file type, "
                     "encryption) before any model is called, and turns it into a clear action for a person.",
            "show": ["Needs action list: the case appears under Blocked.",
                     "Guidance: 'the draft BL is password-protected. SDOC did not open this file or send it to a model.'",
                     "Prepare message: a draft asking for an unprotected copy."],
            "say": "Untrusted attachments are checked before they cost anything or reach a model.",
            "files": lambda d: [
                pdf(d, f"{ref(5)}_SI.pdf", si_doc(ref(5), s5)),
                pdf(d, f"{ref(5)}_BL.pdf", bl_doc(ref(5), s5), password=f"{batch}2605"),
            ],
            "emails": [{
                "subject": f"Draft BL (password protected) - {ref(5)}",
                "body": f"Hi team,\n\nPlease compare the SI and draft BL for {ref(5)}. The carrier sends drafts "
                        f"password protected; the password is {batch}2605.\n\nRegards,\nSiti\nShipping Desk",
                "attach": [f"{ref(5)}_SI.pdf", f"{ref(5)}_BL.pdf"],
                "expect": "BLOCKED",
            }],
        },
        {
            "key": "06-inbox-triage", "ref": ref(6),
            "title": "Not every email is a BL check",
            "result": "Three emails routed as Invoice query, SI request and Spam. No shipment cases created",
            "story": "A shared documentation inbox mostly receives other things: charge queries, new booking "
                     "requests, phishing. All three of these carry attachments, so a tool that treats every "
                     "attachment as work would open three bogus cases. SDOC routes them and keeps the case "
                     "queue clean.",
            "show": ["Routed mail view: three new rows with their categories.",
                     "Needs action and All cases: unchanged."],
            "say": "It reads intent, not just attachments.",
            "files": lambda d: [
                pdf(d, "INV-88214.pdf", ("TAX INVOICE", f"{CARRIER} - local charges, Port Klang", [
                    ("Invoice", [("Invoice No.", "INV-88214"), ("Booking No.", ref(6))]),
                    ("Charges", [("Terminal Handling", "MYR 1,450.00"), ("Documentation Fee", "MYR 180.00"),
                                 ("Seal Fee", "MYR 45.00"), ("Total Due", "MYR 1,675.00")]),
                ])),
                docx_file(d, f"Booking_Confirmation_{ref(6)}.docx", ref(6)),
                pdf(d, "Secure_Document.pdf", ("SECURE MESSAGE", "Mailbox administration", [
                    ("Action", [("Status", "Mailbox storage full"), ("Required", "Verify your account today")]),
                ])),
            ],
            "emails": [
                {
                    "subject": "Invoice INV-88214 - local charges query",
                    "body": "Hi,\n\nThe THC on the attached invoice looks higher than quoted. Can you check "
                            "the local charges before we pay?\n\nThanks,\nAccounts",
                    "attach": ["INV-88214.pdf"],
                    "expect": "INVOICE_QUERY",
                },
                {
                    "subject": f"SI request - new booking {ref(6)}",
                    "body": "Hi team,\n\nBooking confirmed with the carrier; confirmation attached. Please "
                            "prepare the SI for this shipment.\n\nThanks,\nSales Support",
                    "attach": [f"Booking_Confirmation_{ref(6)}.docx"],
                    "expect": "SI_REQUEST",
                },
                {
                    "subject": "Verify account - your mailbox storage is full",
                    "body": "Your mailbox has exceeded its storage limit. Open the attached secure document "
                            "and verify your account within 24 hours to avoid suspension.",
                    "attach": ["Secure_Document.pdf"],
                    "expect": "SPAM",
                },
            ],
        },
    ]


def pdf(folder, name, doc, password=None):
    write_pdf(folder / name, *doc, password=password)
    return name


def xlsx(folder, name, ref, rows):
    write_packing_list(folder / name, ref, rows)
    return name


def docx_file(folder, name, ref):
    write_booking_docx(folder / name, ref)
    return name


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------
def email_text(email):
    return (
        "TO:       the mailbox connected to SDOC (send to ONE mailbox only)\n"
        f"SUBJECT:  {email['subject']}\n"
        f"ATTACH:   {', '.join(email['attach'])}\n"
        f"EXPECT:   {email['expect']}\n"
        "--------------------------------------------------------------------\n"
        f"{email['body']}\n"
    )


def build(out, batch):
    out = Path(out)
    if out.exists():
        shutil.rmtree(out)
    offline = out / "_offline"
    (offline / "inbox").mkdir(parents=True)
    (offline / "attachments").mkdir()
    items = scenarios(batch)
    for index, item in enumerate(items, start=1):
        folder = out / item["key"]
        folder.mkdir(parents=True)
        item["file_names"] = item["files"](folder)
        many = len(item["emails"]) > 1
        for step, email in enumerate(item["emails"], start=1):
            name = f"email-{step}.txt" if many else "email.txt"
            (folder / name).write_text(email_text(email), encoding="utf-8")
            # Machine copy for --check and the offline fallback. Ids sort in send order.
            email_id = f"kit{batch}_{index:02d}{chr(96 + step)}"
            # One folder per email, like the mailbox sync, so file names stay clean.
            (offline / "attachments" / email_id).mkdir()
            for attachment in email["attach"]:
                shutil.copy(folder / attachment, offline / "attachments" / email_id / attachment)
            record = {
                "email_id": email_id,
                "from": "demo.sender@example.com",
                "subject": email["subject"],
                "body": email["body"],
                "attachments": [f"attachments/{email_id}/{a}" for a in email["attach"]],
                "expect": email["expect"],
            }
            (offline / "inbox" / f"{email_id}.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
    (out / "README.md").write_text(runbook(items, batch), encoding="utf-8")
    return items


def runbook(items, batch):
    lines = [
        f"# SDOC live demo kit (batch {batch})",
        "",
        "Six scenarios. Each folder holds the documents to attach and the email text to send.",
        f"Shipment references in this batch: {items[0]['ref']} to {items[-1]['ref']}.",
        "",
        "## Before you present",
        "",
        "1. **Use a fresh batch for every run-through.** A reference that already has a case joins "
        "that case. Rehearse with `--batch 2`, `--batch 3`, and present with a batch you have never sent:",
        "   `python Averis1am/scripts/generate_demo_kit.py --batch 4 --check`",
        "2. **Send to one mailbox only** (the default Outlook mailbox). Sending to both creates duplicates.",
        "3. **No signature images.** Outlook passes signature logos through as attachments, and an extra "
        "image breaks the SI + BL pair. Send plain emails.",
        "4. **Attach the files as ordinary attachments**, not as links or inline images.",
        "5. **Keep the subjects exactly as written.** Words like 'invoice', 'freight', 'THC' or "
        "'reminder' in the subject route an email away from comparison.",
        "6. After sending, press **Sync now** (or wait up to a minute). Each comparison makes two "
        "Gemini reads, so allow 10 to 30 seconds per email.",
        "7. For two-email scenarios, wait until the first result is on screen before sending the second.",
        "",
        "**If the mailbox fails on the day**, load the same scenarios without email. Use a batch you "
        "have not sent, or these cases join the live ones. From the repository root:",
        "",
        "```powershell",
        "Copy-Item -Recurse -Force demo-kit/_offline/attachments/* Averis1am/data/attachments/",
        "cd Averis1am; python scripts/ingest_source.py ../demo-kit/_offline --verify",
        "```",
        "",
        "The copy step lets the evidence pane render the pages; without it the crops show "
        "'page could not be rendered'.",
        "",
        "## Scenarios",
        "",
    ]
    for index, item in enumerate(items, start=1):
        lines += [
            f"### {index}. {item['title']} - `{item['key']}/`",
            "",
            f"**Result:** {item['result']}",
            "",
            item["story"],
            "",
            "**Send:** " + (" then ".join(
                f"`email-{n}.txt` ({e['expect']})" for n, e in enumerate(item["emails"], start=1))
                if len(item["emails"]) > 1 else "`email.txt`"),
            "",
            "**Show:**",
            *[f"- {point}" for point in item["show"]],
            "",
            f"**Say:** \"{item['say']}\"",
            "",
        ]
    lines += [
        "## Honest limits (for questions)",
        "",
        "- Container **count** is compared, container **type** is not: '2 x 40'HC' and '2 x 20'GP' match. "
        "Type can be added as a configured field.",
        "- Each email is compared on its own attachments. A draft BL sent alone after an SI-only email "
        "keeps the case Waiting, which is why scenario 4's second email carries both files.",
        "- The Gemini second read accepts PDFs and images, so the SI and BL must be PDFs for a live "
        "result. Word and Excel files are read as supporting documents.",
        "- SDOC drafts messages but never sends them. A person reviews and sends.",
        "- All companies, vessels and values in this kit are fictitious.",
        "",
    ]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Check: run every email through the same service the mailbox sync uses.
# ---------------------------------------------------------------------------
def check(out, verify=False):
    from dotenv import load_dotenv
    from sdoc.casework import CaseService, CaseStore, shipment_reference
    from sdoc.sources import DirSource

    load_dotenv(ROOT.parent / ".env")
    cfg = {"extractor": "deterministic", "reader_isolation": False}
    verifier = None
    if verify:
        from sdoc.verification import GeminiVerifier
        verifier = cfg["verifier"] = GeminiVerifier()
    source = DirSource(Path(out) / "_offline")
    failures = 0
    try:
        with CaseStore(":memory:") as store:
            service = CaseService(store, cfg)
            for email in source.emails():
                result = service.ingest_email(email, source)
                record, case = result["record"], result["case"]
                got = case["state"] if case else record["category"]
                detail = case["state_reason"] if case else ""
                defects = ",".join(record.get("defect_fields") or [])
                verdict = (result["evidence"].get("verification") or {}).get("status", "-")
                ok = got == email["expect"]
                failures += not ok
                print(f"{'PASS' if ok else 'FAIL'}  {email['email_id']:<12} "
                      f"{shipment_reference(email):<18} expect {email['expect']:<13} "
                      f"got {got:<13} {detail or ''} {defects} verifier={verdict}")
    finally:
        if verifier:
            verifier.close()
    print("all scenarios behave as scripted" if not failures else f"{failures} email(s) off script")
    return failures


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--batch", type=int, default=1, choices=range(1, 10),
                        help="1-9; each batch gets its own shipment references")
    parser.add_argument("--out", default=str(ROOT.parent / "demo-kit"))
    parser.add_argument("--check", action="store_true", help="run every email through the pipeline")
    parser.add_argument("--verify", action="store_true", help="include the Gemini second read in --check")
    args = parser.parse_args()
    items = build(args.out, args.batch)
    print(f"Wrote {len(items)} scenarios to {args.out} (batch {args.batch})")
    if args.check:
        sys.exit(1 if check(args.out, args.verify) else 0)


if __name__ == "__main__":
    main()
