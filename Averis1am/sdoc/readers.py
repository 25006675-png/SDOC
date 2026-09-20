"""Attachment readers.

Every reader takes raw bytes and returns ``(title, [(label, value), ...])``
or ``None`` when the file cannot be read at all (-> 'unreadable').

Readers are looked up by file extension in ``READERS`` — new formats plug in
via ``register_reader('.ext', fn)`` without touching the pipeline.
"""
import io
import re


def _cfg_flag(cfg, name, default=False):
    """Read a feature flag from the pipeline's dict config or an object."""
    if isinstance(cfg, dict):
        return cfg.get(name, default)
    return getattr(cfg, name, default)

# ---------------------------------------------------------------------------
# Optional OCR (scanned PDFs / image attachments). Enabled per-run via cfg.ocr
# and only used when the native text layer is absent.
# ---------------------------------------------------------------------------
def ocr_available():
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
        return bool(pytesseract.get_tesseract_version())
    except Exception:
        return False


def _ocr_pdf(data):
    import pypdfium2 as pdfium
    import pytesseract
    pdf = pdfium.PdfDocument(data)
    return "\n".join(
        pytesseract.image_to_string(page.render(scale=2).to_pil())
        for page in pdf)


def _ocr_image(data):
    import pytesseract
    from PIL import Image
    return pytesseract.image_to_string(Image.open(io.BytesIO(data)))


def _pairs_from_ocr(text):
    """Parse OCR'd plain text like a .txt document."""
    if not text or not text.strip():
        return None
    return _txt_pairs(text.splitlines())


# ---------------------------------------------------------------------------
def _txt_pairs(lines):
    title = ""
    for l in lines:
        s = l.strip()
        if s and not set(s) <= set("=-_*#"):
            title = s
            break
    pairs = []
    for l in lines:
        if not l.strip() or l[0] in " \t":      # indented lines are continuations
            continue
        if ":" in l:
            lab, val = l.split(":", 1)
            pairs.append((lab.strip(), val.strip()))
    return title, pairs


def read_txt(data, cfg=None):
    # try encodings in order: utf-8 -> CJK legacy -> latin-1 (never fails)
    for enc in ("utf-8", "gb18030", "big5", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        return None
    return _txt_pairs(text.splitlines())


# ---------------------------------------------------------------------------
def _join_chars(cs):
    """Rebuild text from pdf chars, inserting spaces at >1.5pt gaps."""
    out, prev = "", None
    for c in cs:
        if prev is not None and c["x0"] - prev["x1"] > 1.5:
            out += " "
        out += c["text"]
        prev = c
    return out


def _pairs_from_pdf_chars(chars):
    """Rows clustered by vertical overlap; each row split by font — labels are
    bold, values regular. Robust to labels overlapping the value column and
    to odd font metrics (CJK) that break geometric column splitting."""
    chars.sort(key=lambda c: (c["top"], c["x0"]))
    rows = []
    for c in chars:
        if rows and c["top"] <= rows[-1][1] + 1.5:
            rows[-1][0].append(c)
            rows[-1][1] = max(rows[-1][1], c["bottom"])
        else:
            rows.append([[c], c["bottom"]])
    lines = []
    for rchars, _ in rows:
        rchars.sort(key=lambda c: c["x0"])
        bold = _join_chars([c for c in rchars if "bold" in c["fontname"].lower()])
        reg = _join_chars([c for c in rchars if "bold" not in c["fontname"].lower()])
        lines.append((bold.strip(), reg.strip(), _join_chars(rchars).strip()))
    title = next((t for _, _, t in lines if t), "")
    pairs = []
    for bold, reg, _ in lines:
        if ":" in bold:                          # "Label: value" in one draw run
            lab, val = bold.split(":", 1)
            pairs.append((lab.strip(), f"{val} {reg}".strip()))
        elif bold and reg:
            pairs.append((bold, reg))            # bold label col + regular value
    return title, pairs


def read_pdf(data, cfg=None):
    import pdfplumber
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            chars = []
            for idx, pg in enumerate(pdf.pages, start=1):
                for c in pg.chars:
                    item = dict(c)
                    item["page_number"] = getattr(pg, "page_number", idx) or idx
                    chars.append(item)
    except Exception:
        return None                                # corrupt / truncated file
    if chars:
        return _pairs_from_pdf_chars(chars)
    # no text layer: scanned copy — OCR if enabled, else unreadable
    if _cfg_flag(cfg, "ocr") and ocr_available():
        try:
            return _pairs_from_ocr(_ocr_pdf(data))
        except Exception:
            return None
    return None


# ---------------------------------------------------------------------------
def read_docx(data, cfg=None):
    import docx
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        return None
    title = next((p.text.strip() for p in d.paragraphs if p.text.strip()), "")
    pairs = []
    for t in d.tables:
        for row in t.rows:
            if len(row.cells) >= 2:
                lab, val = row.cells[0].text.strip(), row.cells[1].text.strip()
                if lab and val:
                    pairs.append((lab, val))
    return title, pairs


def read_xlsx(data, cfg=None):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    except Exception:
        return None
    ws = wb.active
    pairs = []
    for row in ws.iter_rows(values_only=True):
        if len(row) >= 2 and row[0] is not None and row[1] is not None:
            pairs.append((str(row[0]).strip(), str(row[1]).strip()))
    return (ws.title or ""), pairs


def read_image(data, cfg=None):
    """Image attachments are only readable via OCR."""
    if _cfg_flag(cfg, "ocr") and ocr_available():
        try:
            return _pairs_from_ocr(_ocr_image(data))
        except Exception:
            return None
    return None


def read_unsupported(data, cfg=None):
    return None


# ---------------------------------------------------------------------------
# Extension registry — plug new formats in with register_reader().
# ---------------------------------------------------------------------------
READERS = {
    ".txt": read_txt, ".text": read_txt, ".csv": read_txt, ".md": read_txt,
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".xlsx": read_xlsx, ".xlsm": read_xlsx,
    ".png": read_image, ".jpg": read_image, ".jpeg": read_image,
    ".tif": read_image, ".tiff": read_image, ".bmp": read_image,
    ".doc": read_unsupported,                    # legacy binary Word
}


def register_reader(ext, fn):
    READERS[ext.lower()] = fn


def reader_for(path):
    """Reader function for a file path's extension, or None."""
    name = str(path).lower()
    ext = "." + name.rsplit(".", 1)[-1] if "." in name else ""
    return READERS.get(ext)
